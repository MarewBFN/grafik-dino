"""Karta pracy pojedynczego pracownika - jedna strona A4 = jeden
pracownik = jeden miesiąc, wzorowana na papierowej "Liście obecności
miesięcznej pracownika" klienta. Zastępuje dawny eksport "(jeden
pracownik)" (ten sam co eksport całej załogi, tylko z jedną osobą) -
to jest osobny, od zera zbudowany układ: dni w WIERSZACH, nie w
kolumnach, z miejscem na odręczny podpis. Niezależny od
`export/image_exporter.py`/`export/excel_exporter.py` i od
`export/security_*_exporter.py` - uniwersalny dla każdego profilu
działalności, nie warunkowany business_type.

"Norma" (nominalny wymiar godzin na miesiąc) na razie świadomie zostaje
pusta, do potwierdzenia z klientem - patrz "plan profil ochrona (analiza
specyfikacji klienta).md", sekcja 16.

Kolumny "Godziny dzienne"/"Godziny nocne" liczą, ile z godzin danej zmiany
przypada w oknie 22:00-06:00 (pora nocna wg kodeksu pracy) a ile poza nim -
patrz _night_minutes()/_day_night_hours(). Karta ma tylko jedną komórkę na
podpis, na dole strony (nie osobną kolumnę na każdy dzień).
"""

import calendar
import re
from datetime import datetime

from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_TITLE = "Lista obecności miesięczna pracownika"
_COLUMNS = ["Dzień", "Wejście", "Wyjście", "Ilość godzin", "Godziny dzienne", "Godziny nocne"]

_NIGHT_WINDOW_START = 22 * 60  # 22:00 w minutach od północy
_NIGHT_WINDOW_END = 6 * 60  # 06:00 w minutach od północy (następnej doby)
_MINUTES_PER_DAY = 24 * 60


def _minutes(time_str: str) -> int:
    h, m = time_str.split(":")
    return int(h) * 60 + int(m)


def _overlap(a_start: int, a_end: int, b_start: int, b_end: int) -> int:
    return max(0, min(a_end, b_end) - max(a_start, b_start))


def _night_minutes(ds) -> int:
    """Minuty zmiany danego dnia leżące w oknie 22:00-06:00 (definicja
    "pory nocnej" z kodeksu pracy). Zmiana 24h zawiera to okno w całości
    dokładnie raz."""
    if ds.is_empty() or ds.is_leave or getattr(ds, "is_sick", False):
        return 0

    if ds.is_full_day:
        return _MINUTES_PER_DAY - _NIGHT_WINDOW_START + _NIGHT_WINDOW_END  # 8h

    start_abs = _minutes(ds.start)
    duration = ds.total_duration()
    if duration is None:
        return 0
    end_abs = start_abs + int(duration.total_seconds() // 60)

    night = 0
    for k in (-1, 0, 1):
        window_start = _NIGHT_WINDOW_START + k * _MINUTES_PER_DAY
        window_end = _MINUTES_PER_DAY + _NIGHT_WINDOW_END + k * _MINUTES_PER_DAY
        night += _overlap(start_abs, end_abs, window_start, window_end)

    return night


def _format_minutes(total_minutes: int) -> str:
    hours, minutes = divmod(total_minutes, 60)
    return f"{hours}:{minutes:02d}"


def _day_night_hours(ds) -> tuple[str, str]:
    """(godziny dzienne, godziny nocne) danego dnia, jako stringi "H:MM" -
    puste dla dni bez zmiany/urlopu/L4 (kolumny mają być wtedy puste,
    zgodnie z resztą wiersza)."""
    if ds.is_empty() or ds.is_leave or getattr(ds, "is_sick", False):
        return "", ""

    duration = ds.total_duration()
    if duration is None:
        return "", ""

    total = int(duration.total_seconds() // 60)
    night = _night_minutes(ds)
    day = total - night
    return _format_minutes(day), _format_minutes(night)


def _format_hour(time_str):
    """Zawsze pełny zapis "H:MM" (np. "8:00", nie "8") - godzina bez
    zera wiodącego, minuty zawsze dwucyfrowe, spójnie z kolumnami
    "Godziny dzienne"/"Godziny nocne" (patrz _format_minutes)."""
    if not time_str:
        return ""
    h, m = time_str.split(":")
    return f"{int(h)}:{m}"


def _location_name(shop, employee) -> str:
    """Nazwa lokalizacji/placówki przypisanej pracownikowi - to jest
    wartość pola "Stanowisko" na karcie (ustalone z użytkownikiem: w tej
    branży stanowisko pracownika to jego posterunek/placówka, nie
    osobne pole kadrowe)."""
    if shop is None:
        return ""
    location = getattr(shop, "locations", {}).get(employee.location_key)
    if location is not None and getattr(location, "name", ""):
        return location.name
    return getattr(shop, "name", "") or ""


def _day_rows(schedule, employee):
    """(dzień, wejście, wyjście, ilość_godzin, godziny_dzienne,
    godziny_nocne) dla każdego dnia miesiąca. Urlop/L4 dostają jawny
    znacznik w kolumnie "Wejście" ("Urlop"/"L4") zamiast być nieodróżnialne
    od dnia bez żadnej zmiany (dawniej wszystkie trzy przypadki dawały
    identyczny pusty wiersz) - dzień bez zmiany (ani urlopu, ani L4) dalej
    zostaje pusty. Koniec zmiany przechodzącej przez północ pokazuje samą
    godzinę, bez znacznika "+1" (usunięty na życzenie użytkownika - karta
    ma jeden podpis na cały miesiąc, nie osobną kolumnę na dzień, więc
    ujednoznacznienie dnia i tak nie ma tu zastosowania jak w siatce
    grafiku)."""
    rows = []
    for day in range(1, schedule.days_in_month + 1):
        ds = schedule.get_day(employee, day)
        if ds.is_leave:
            rows.append((day, "Urlop", "", "", "", ""))
            continue
        if getattr(ds, "is_sick", False):
            rows.append((day, "L4", "", "", "", ""))
            continue
        if ds.is_empty():
            rows.append((day, "", "", "", "", ""))
            continue

        day_hours, night_hours = _day_night_hours(ds)
        rows.append((
            day, _format_hour(ds.start), _format_hour(ds.end), ds.total_as_str() or "",
            day_hours, night_hours,
        ))

    return rows


def sanitize_filename_part(text: str) -> str:
    """Usuwa znaki niedozwolone w nazwach plików Windows - używane też
    przez `ui/main_window.py` do budowania nazw plików przy eksporcie
    zbiorczym (jeden JPG per pracownik w folderze)."""
    return re.sub(r'[\\/:*?"<>|]', "_", text).strip()


# ============================= IMAGE =============================

_PAGE_W, _PAGE_H = 1240, 1754  # A4 @ 150dpi, portret


class _EmployeeCardImageExporter:
    def __init__(self, schedule, year, month, shop, employee):
        self.schedule = schedule
        self.year = year
        self.month = month
        self.shop = shop
        self.employee = employee

        self.MARGIN = 60
        self.img = Image.new("RGB", (_PAGE_W, _PAGE_H), "white")
        self.draw = ImageDraw.Draw(self.img)

        try:
            self.font = ImageFont.truetype("arial.ttf", 16)
            self.font_b = ImageFont.truetype("arial.ttf", 22)
        except Exception:
            self.font = ImageFont.load_default()
            self.font_b = self.font

        self.BLACK = (0, 0, 0)
        self.TITLE_BG = (237, 237, 237)

    def render(self):
        """Rysuje kartę i zwraca gotowy obraz - bez zapisu, żeby ten sam
        render mógł posłużyć zarówno JPG (`export`), jak i PDF
        (`export_employee_cards_to_pdf`, wielostronicowy - jedna strona
        na pracownika) bez powielania logiki rysowania."""
        y = self._draw_header()
        y = self._draw_data_block(y)
        self._draw_table(y)
        return self.img

    def export(self, path):
        self.render()
        self.img.save(path, "JPEG", quality=95)

    def _draw_header(self):
        y = self.MARGIN
        self.draw.text((self.MARGIN, y), _TITLE, fill=self.BLACK, font=self.font_b)
        y += 40

        self._draw_dashed_line(self.MARGIN, y, _PAGE_W - self.MARGIN, y)
        return y + 30

    def _draw_dashed_line(self, x0, y0, x1, y1, dash=8, gap=6):
        x = x0
        while x < x1:
            x_end = min(x + dash, x1)
            self.draw.line([(x, y0), (x_end, y1)], fill=self.BLACK)
            x += dash + gap

    def _draw_data_block(self, y):
        block_h = 3 * 34
        left_w = int((_PAGE_W - 2 * self.MARGIN) * 0.32)
        left_x = self.MARGIN
        right_x = self.MARGIN + left_w + 20
        right_w = (_PAGE_W - self.MARGIN) - right_x

        left_rows = [
            ("Rok", str(self.year)),
            ("M-c", f"{self.month:02d}"),
            ("Norma", ""),
        ]
        row_h = block_h // 3
        for i, (label, value) in enumerate(left_rows):
            ry = y + i * row_h
            value_x = left_x + left_w // 2
            self.draw.rectangle([left_x, ry, value_x, ry + row_h], fill=self.TITLE_BG)
            self.draw.rectangle([left_x, ry, left_x + left_w, ry + row_h], outline=self.BLACK)
            self.draw.text((left_x + 6, ry + row_h // 2), label, fill=self.BLACK, font=self.font, anchor="lm")
            self.draw.line([(value_x, ry), (value_x, ry + row_h)], fill=self.BLACK)
            self.draw.text((value_x + 6, ry + row_h // 2), value, fill=self.BLACK, font=self.font, anchor="lm")

        right_row_h = block_h // 2
        right_label_h = 22
        right_rows = [
            ("Imię i nazwisko pracownika", self.employee.display_name()),
            ("Stanowisko", _location_name(self.shop, self.employee)),
        ]
        for i, (label, value) in enumerate(right_rows):
            ry = y + i * right_row_h
            self.draw.rectangle([right_x, ry, right_x + right_w, ry + right_label_h], fill=self.TITLE_BG)
            self.draw.rectangle([right_x, ry, right_x + right_w, ry + right_row_h], outline=self.BLACK)
            self.draw.text((right_x + 6, ry + 4), label, fill=self.BLACK, font=self.font, anchor="la")
            self.draw.text((right_x + 6, ry + right_row_h - 6), value, fill=self.BLACK, font=self.font_b, anchor="lb")

        return y + block_h + 30

    def _column_widths(self):
        """Szerokości 6 kolumn tabeli (patrz _COLUMNS) - "Godziny dzienne"/
        "Godziny nocne" dostają dokładnie tyle samo miejsca (cała
        szerokość zostająca po pierwszych czterech kolumnach, podzielona
        na pół); dawniej ostatnia kolumna brała całą resztę, więc
        "nocne" wychodziło ponad 2x szersze niż "dzienne", mimo że mają
        identyczną treść (format "H:MM")."""
        table_w = (_PAGE_W - self.MARGIN) - self.MARGIN
        fixed_w = [70, 150, 150, 150]
        hours_w = table_w - sum(fixed_w)
        night_w = hours_w // 2
        day_w = hours_w - night_w
        return fixed_w + [day_w, night_w]

    def _draw_table(self, y):
        table_x0 = self.MARGIN
        col_w = self._column_widths()
        col_x = [table_x0]
        for w in col_w:
            col_x.append(col_x[-1] + w)

        header_h = 36
        for i, header in enumerate(_COLUMNS):
            self.draw.rectangle([col_x[i], y, col_x[i + 1], y + header_h], fill=self.TITLE_BG, outline=self.BLACK)
            self._centered_text((col_x[i] + col_x[i + 1]) // 2, y + header_h // 2, header, self.font)
        y += header_h

        row_h = 24
        rows = _day_rows(self.schedule, self.employee)
        for day, start, end, hours, day_hours, night_hours in rows:
            values = [str(day), start, end, hours, day_hours, night_hours]
            for i, val in enumerate(values):
                self.draw.rectangle([col_x[i], y, col_x[i + 1], y + row_h], outline=self.BLACK)
                if val:
                    self._centered_text((col_x[i] + col_x[i + 1]) // 2, y + row_h // 2, val, self.font)
            y += row_h

        total = self.schedule.total_hours_for_employee(self.employee)
        footer_h = 30
        self.draw.rectangle([col_x[0], y, col_x[3], y + footer_h], fill=self.TITLE_BG, outline=self.BLACK)
        self.draw.text((col_x[0] + 6, y + footer_h // 2), "Razem ilość godzin:", fill=self.BLACK, font=self.font, anchor="lm")
        self.draw.rectangle([col_x[3], y, col_x[4], y + footer_h], outline=self.BLACK)
        self._centered_text((col_x[3] + col_x[4]) // 2, y + footer_h // 2, str(total), self.font_b)
        # Jedyna komórka na podpis na całej karcie (bez kolumny na podpis w
        # każdym dniu) - podpisana, bo bez nagłówka kolumny nie byłoby wiadomo,
        # co to za puste pole.
        self.draw.rectangle([col_x[4], y, col_x[6], y + footer_h], outline=self.BLACK)
        self.draw.text(
            (col_x[4] + 6, y + footer_h // 2), "Podpis pracownika:",
            fill=self.BLACK, font=self.font, anchor="lm",
        )

        return y + footer_h

    def _centered_text(self, x, y, text, font):
        bbox = self.draw.textbbox((0, 0), text, font=font)
        w = bbox[2] - bbox[0]
        h = bbox[3] - bbox[1]
        self.draw.text((x - w // 2, y - h // 2), text, fill=self.BLACK, font=font)


def render_employee_card_image(schedule, year, month, shop, employee):
    """Renderuje kartę pracy jednego pracownika do obrazu (PIL Image),
    bez zapisu do pliku - współdzielone przez `export_employee_card_to_image`/
    `export_employee_cards_to_pdf` oraz przez podgląd przed eksportem
    (`ui/export_preview_dialog.py`), żeby podgląd był dokładnie tym, co
    trafi do pliku."""
    return _EmployeeCardImageExporter(schedule, year, month, shop, employee).render()


def export_employee_card_to_image(schedule, year, month, path, shop=None, employee=None):
    image = render_employee_card_image(schedule, year, month, shop, employee)
    image.save(path, "JPEG", quality=95)
    return True


# ============================== PDF ===============================


def save_employee_card_pages_to_pdf(pages, path):
    """Zapisuje gotowe obrazy kart (jedna strona = jeden pracownik, patrz
    `render_employee_card_image`) jako jeden wielostronicowy PDF (Pillow
    wspiera to natywnie przez `save_all`/`append_images`)."""
    first, rest = pages[0], pages[1:]
    first.save(path, "PDF", resolution=150.0, save_all=True, append_images=rest)


def export_employee_cards_to_pdf(schedule, year, month, path, shop=None, employees=None):
    """Jeden dokument PDF, jedna strona na pracownika (kolejność jak w
    `employees`) - ten sam render co JPG (`render_employee_card_image`),
    tylko zapisany jako PDF zamiast JPEG."""
    employees = employees if employees is not None else schedule.employees
    pages = [render_employee_card_image(schedule, year, month, shop, employee) for employee in employees]
    save_employee_card_pages_to_pdf(pages, path)
    return True


# ============================= EXCEL =============================

_DASHED_BOTTOM = Border(bottom=Side(style="dashed"))
_THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
_FILL_TITLE = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")


def _write_employee_sheet(ws, schedule, year, month, shop, employee):
    ws.cell(row=1, column=1, value=_TITLE).font = Font(bold=True, size=14)

    last_col = len(_COLUMNS)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    for c in range(1, last_col + 1):
        ws.cell(row=2, column=c).border = _DASHED_BOTTOM

    ws.cell(row=3, column=1, value="Rok").font = Font(bold=True)
    ws.cell(row=3, column=1).fill = _FILL_TITLE
    ws.cell(row=3, column=2, value=year)
    ws.cell(row=4, column=1, value="M-c").font = Font(bold=True)
    ws.cell(row=4, column=1).fill = _FILL_TITLE
    ws.cell(row=4, column=2, value=f"{month:02d}")
    ws.cell(row=5, column=1, value="Norma").font = Font(bold=True)
    ws.cell(row=5, column=1).fill = _FILL_TITLE

    ws.cell(row=3, column=3, value="Imię i nazwisko pracownika").font = Font(bold=True)
    ws.cell(row=3, column=3).fill = _FILL_TITLE
    ws.merge_cells(start_row=3, start_column=4, end_row=3, end_column=last_col)
    ws.cell(row=3, column=4, value=employee.display_name())

    ws.cell(row=4, column=3, value="Stanowisko").font = Font(bold=True)
    ws.cell(row=4, column=3).fill = _FILL_TITLE
    ws.merge_cells(start_row=4, start_column=4, end_row=4, end_column=last_col)
    ws.cell(row=4, column=4, value=_location_name(shop, employee))

    header_row = 7
    for i, header in enumerate(_COLUMNS):
        cell = ws.cell(row=header_row, column=i + 1, value=header)
        cell.font = Font(bold=True)
        cell.alignment = _ALIGN_CENTER
        cell.border = _THIN_BORDER
        cell.fill = _FILL_TITLE

    row = header_row + 1
    for day, start, end, hours, day_hours, night_hours in _day_rows(schedule, employee):
        values = [day, start, end, hours, day_hours, night_hours]
        for i, val in enumerate(values):
            cell = ws.cell(row=row, column=i + 1, value=val if val != "" else None)
            cell.alignment = _ALIGN_CENTER
            cell.border = _THIN_BORDER
        row += 1

    total = schedule.total_hours_for_employee(employee)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    footer_label = ws.cell(row=row, column=1, value="Razem ilość godzin:")
    footer_label.font = Font(bold=True)
    for c in range(1, 4):
        ws.cell(row=row, column=c).fill = _FILL_TITLE
    total_cell = ws.cell(row=row, column=4, value=total)
    total_cell.font = Font(bold=True)
    total_cell.alignment = _ALIGN_CENTER
    # Jedyna komórka na podpis na całym arkuszu (bez kolumny na podpis w
    # każdym dniu) - w wierszu stopki, pod kolumnami godzin dziennych/nocnych.
    ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=last_col)
    signature_cell = ws.cell(row=row, column=5, value="Podpis pracownika:")
    signature_cell.font = Font(bold=True)
    for c in range(1, last_col + 1):
        ws.cell(row=row, column=c).border = _THIN_BORDER

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 22


def export_employee_cards_to_excel(schedule, year, month, path, shop=None, employees=None):
    employees = employees if employees is not None else schedule.employees

    wb = Workbook()
    wb.remove(wb.active)

    used_titles = set()
    for employee in employees:
        title = sanitize_filename_part(employee.display_name())[:31] or "Pracownik"
        base_title = title
        n = 2
        while title in used_titles:
            suffix = f" ({n})"
            title = base_title[: 31 - len(suffix)] + suffix
            n += 1
        used_titles.add(title)

        ws = wb.create_sheet(title=title)
        _write_employee_sheet(ws, schedule, year, month, shop, employee)

    wb.save(path)
    return True
