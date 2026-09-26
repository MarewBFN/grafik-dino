from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from export.employee_card_exporter import (
    _day_night_hours,
    _day_night_totals,
    _day_rows,
    _EmployeeCardImageExporter,
    _format_hour,
    _night_minutes,
    _nominal_hours_str,
    _overtime_str,
    _strip_zero_minutes,
    export_employee_card_to_image,
    export_employee_cards_to_excel,
)
from model.day_schedule import DaySchedule
from model.employee import Employee
from model.month_schedule import MonthSchedule
from model.shop_config import ShopConfig, DEFAULT_LOCATION_KEY


def _row_for_day(schedule, employee, day):
    prefix = f"{day} "
    return next(r for r in _day_rows(schedule, employee) if r[0].startswith(prefix))


class DayRowsMarkerTests(unittest.TestCase):
    """_day_rows() - urlop/L4 dostają jawny znacznik w kolumnie "Wejście"
    zamiast być nieodróżnialne od dnia bez żadnej zmiany (dawniej
    wszystkie trzy przypadki dawały identyczny pusty wiersz). Koniec
    zmiany przez północ NIE dostaje już znacznika "+1" (usunięty na
    życzenie użytkownika). Dzień 3 sierpnia 2026 to poniedziałek."""

    def _schedule_with_one_employee(self):
        schedule = MonthSchedule(2026, 8)
        emp = Employee(last_name="Kowalski", first_name="Adam")
        schedule.add_employee(emp)
        return schedule, emp

    def test_day_number_has_weekday_abbreviation_attached(self):
        schedule, emp = self._schedule_with_one_employee()

        row = _row_for_day(schedule, emp, 3)

        self.assertEqual(row[0], "3 Pn")

    def test_leave_day_shows_urlop_in_wejscie_column(self):
        schedule, emp = self._schedule_with_one_employee()
        schedule.get_day(emp, 3).set_leave()

        row = _row_for_day(schedule, emp, 3)

        self.assertEqual(row, ("3 Pn", "Urlop", "", "", "", ""))

    def test_sick_day_shows_l4_in_wejscie_column(self):
        schedule, emp = self._schedule_with_one_employee()
        schedule.get_day(emp, 3).set_sick()

        row = _row_for_day(schedule, emp, 3)

        self.assertEqual(row, ("3 Pn", "L4", "", "", "", ""))

    def test_day_without_any_shift_stays_blank(self):
        schedule, emp = self._schedule_with_one_employee()

        row = _row_for_day(schedule, emp, 3)

        self.assertEqual(row, ("3 Pn", "", "", "", "", ""))

    def test_shift_crossing_midnight_has_no_plus_one_marker(self):
        schedule, emp = self._schedule_with_one_employee()
        schedule.get_day(emp, 3).set_hours("22:00", "06:00")

        row = _row_for_day(schedule, emp, 3)

        self.assertEqual(row[1], "22:00")
        self.assertEqual(row[2], "6:00")
        self.assertNotIn("+1", row[2])

    def test_full_hour_shift_drops_the_zero_minute_suffix(self):
        schedule, emp = self._schedule_with_one_employee()
        schedule.get_day(emp, 3).set_hours("08:00", "16:00")

        row = _row_for_day(schedule, emp, 3)

        self.assertEqual(row[3], "8")

    def test_shift_with_non_zero_minutes_keeps_hmm_format(self):
        schedule, emp = self._schedule_with_one_employee()
        schedule.get_day(emp, 3).set_hours("08:00", "16:30")

        row = _row_for_day(schedule, emp, 3)

        self.assertEqual(row[3], "8:30")


class NightMinutesTests(unittest.TestCase):
    def test_plain_day_shift_has_no_night_minutes(self):
        ds = DaySchedule()
        ds.set_hours("08:00", "16:00")
        self.assertEqual(_night_minutes(ds), 0)
        self.assertEqual(_day_night_hours(ds), ("8:00", "0:00"))

    def test_classic_night_shift_is_fully_night(self):
        ds = DaySchedule()
        ds.set_hours("22:00", "06:00")
        self.assertEqual(_night_minutes(ds), 480)
        self.assertEqual(_day_night_hours(ds), ("0:00", "8:00"))

    def test_partial_night_overlap(self):
        ds = DaySchedule()
        ds.set_hours("20:00", "23:00")
        self.assertEqual(_night_minutes(ds), 60)
        self.assertEqual(_day_night_hours(ds), ("2:00", "1:00"))

    def test_full_day_shift_counts_as_night(self):
        ds = DaySchedule()
        ds.set_full_day_shift("06:00")
        self.assertEqual(_night_minutes(ds), 480)
        self.assertEqual(_day_night_hours(ds), ("16:00", "8:00"))

    def test_empty_leave_and_sick_days_have_no_hours(self):
        empty = DaySchedule()
        self.assertEqual(_day_night_hours(empty), ("", ""))

        leave = DaySchedule()
        leave.set_leave()
        self.assertEqual(_day_night_hours(leave), ("", ""))

        sick = DaySchedule()
        sick.set_sick()
        self.assertEqual(_day_night_hours(sick), ("", ""))


class FormatHourTests(unittest.TestCase):
    def test_full_hour_keeps_minutes_instead_of_stripping_them(self):
        self.assertEqual(_format_hour("08:00"), "8:00")

    def test_drops_leading_zero_but_not_minutes(self):
        self.assertEqual(_format_hour("16:00"), "16:00")
        self.assertEqual(_format_hour("00:00"), "0:00")

    def test_non_zero_minutes_are_untouched(self):
        self.assertEqual(_format_hour("08:30"), "8:30")

    def test_empty_string_stays_empty(self):
        self.assertEqual(_format_hour(""), "")


class StripZeroMinutesTests(unittest.TestCase):
    def test_whole_hours_drop_the_zero_minute_suffix(self):
        self.assertEqual(_strip_zero_minutes("8:00"), "8")
        self.assertEqual(_strip_zero_minutes("24:00"), "24")

    def test_non_zero_minutes_are_kept(self):
        self.assertEqual(_strip_zero_minutes("8:30"), "8:30")

    def test_empty_string_stays_empty(self):
        self.assertEqual(_strip_zero_minutes(""), "")


class NominalHoursTests(unittest.TestCase):
    def test_none_shop_returns_empty_string(self):
        emp = Employee(last_name="Testowy", first_name="Jan")
        self.assertEqual(_nominal_hours_str(None, emp), "")

    def test_matches_full_time_nominal_hours_for_full_etat(self):
        shop = ShopConfig(2026, 8)
        emp = Employee(last_name="Testowy", first_name="Jan")

        nominal_hours = shop.get_full_time_nominal_hours()
        hours, minutes = divmod(round(nominal_hours * 60), 60)
        self.assertEqual(_nominal_hours_str(shop, emp), f"{hours}:{minutes:02d}")

    def test_scales_with_employment_fraction(self):
        shop = ShopConfig(2026, 8)
        emp = Employee(last_name="Testowy", first_name="Jan", employment_fraction=0.5)

        nominal_hours = shop.get_full_time_nominal_hours() * 0.5
        hours, minutes = divmod(round(nominal_hours * 60), 60)
        self.assertEqual(_nominal_hours_str(shop, emp), f"{hours}:{minutes:02d}")


class OvertimeStrTests(unittest.TestCase):
    def test_none_shop_returns_empty_string(self):
        schedule, emp = MonthSchedule(2026, 8), Employee(last_name="Testowy", first_name="Jan")
        schedule.add_employee(emp)
        self.assertEqual(_overtime_str(schedule, None, emp), "")

    def test_no_overtime_when_under_nominal(self):
        shop = ShopConfig(2026, 8)
        schedule = MonthSchedule(2026, 8)
        emp = Employee(last_name="Testowy", first_name="Jan")
        schedule.add_employee(emp)
        schedule.get_day(emp, 3).set_hours("08:00", "16:00")

        self.assertEqual(_overtime_str(schedule, shop, emp), "0:00")


class DayNightTotalsTests(unittest.TestCase):
    def test_sums_day_and_night_minutes_across_the_month(self):
        schedule = MonthSchedule(2026, 8)
        emp = Employee(last_name="Testowy", first_name="Jan")
        schedule.add_employee(emp)
        schedule.get_day(emp, 1).set_hours("08:00", "16:00")
        schedule.get_day(emp, 2).set_hours("20:00", "23:00")

        day_total, night_total = _day_night_totals(schedule, emp)

        self.assertEqual(day_total, "10:00")
        self.assertEqual(night_total, "1:00")

    def test_leave_and_sick_days_are_ignored(self):
        schedule = MonthSchedule(2026, 8)
        emp = Employee(last_name="Testowy", first_name="Jan")
        schedule.add_employee(emp)
        schedule.get_day(emp, 1).set_leave()
        schedule.get_day(emp, 2).set_sick()

        self.assertEqual(_day_night_totals(schedule, emp), ("0:00", "0:00"))


def _sample_schedule():
    schedule = MonthSchedule(2026, 8)
    emp = Employee(last_name="Kowalski", first_name="Adam", location_key=DEFAULT_LOCATION_KEY)
    schedule.add_employee(emp)
    schedule.get_day(emp, 1).set_hours("08:00", "16:00")
    schedule.get_day(emp, 2).set_leave()
    return schedule, emp


class ExcelCardExportTests(unittest.TestCase):
    def test_single_employee_creates_one_sheet_with_expected_header(self):
        schedule, emp = _sample_schedule()
        shop = ShopConfig(2026, 8)

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/karta.xlsx"
            export_employee_cards_to_excel(schedule, 2026, 8, path, shop=shop, employees=[emp])

            wb = load_workbook(path)
            self.assertEqual(len(wb.sheetnames), 1)
            ws = wb.active

            self.assertEqual(ws.cell(row=1, column=1).value, "Lista obecności miesięczna pracownika")
            self.assertEqual(ws.cell(row=3, column=2).value, 2026)
            self.assertEqual(ws.cell(row=4, column=2).value, "Sierpień")
            self.assertEqual(ws.cell(row=5, column=2).value, _nominal_hours_str(shop, emp))
            self.assertEqual(ws.cell(row=3, column=4).value, emp.display_name())
            # "Stanowisko" zostaje puste - nie zaczytujemy już nazwy placówki.
            self.assertIsNone(ws.cell(row=4, column=4).value)

    def test_norma_is_empty_without_a_shop(self):
        schedule, emp = _sample_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/karta.xlsx"
            export_employee_cards_to_excel(schedule, 2026, 8, path, employees=[emp])

            wb = load_workbook(path)
            ws = wb.active
            self.assertIsNone(ws.cell(row=5, column=2).value)

    def test_leave_day_shows_urlop_marker_and_leaves_the_rest_blank(self):
        schedule, emp = _sample_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/karta.xlsx"
            export_employee_cards_to_excel(schedule, 2026, 8, path, employees=[emp])

            wb = load_workbook(path)
            ws = wb.active
            # Wiersz nagłówka tabeli to 7, dzień 1 to wiersz 8, dzień 2 (urlop, niedziela) to wiersz 9.
            leave_row = 9
            self.assertEqual(ws.cell(row=leave_row, column=1).value, "2 N")
            # Kolumna 2 to "Wejście" - tam ląduje znacznik "Urlop" zamiast
            # nieodróżnialnego pustego wiersza (patrz _day_rows()).
            self.assertEqual(ws.cell(row=leave_row, column=2).value, "Urlop")
            for col in range(3, 7):
                self.assertIsNone(ws.cell(row=leave_row, column=col).value)

    def test_full_hour_shift_drops_the_zero_minute_suffix_in_excel(self):
        schedule, emp = _sample_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/karta.xlsx"
            export_employee_cards_to_excel(schedule, 2026, 8, path, employees=[emp])

            wb = load_workbook(path)
            ws = wb.active
            # Dzień 1 (zmiana 08:00-16:00, 8h równe) to wiersz 8, kolumna 4 "Ilość godzin".
            self.assertEqual(ws.cell(row=8, column=4).value, "8")

    def test_footer_sums_hours_overtime_and_signature_below_table(self):
        schedule, emp = _sample_schedule()
        shop = ShopConfig(2026, 8)

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/karta.xlsx"
            export_employee_cards_to_excel(schedule, 2026, 8, path, shop=shop, employees=[emp])

            wb = load_workbook(path)
            ws = wb.active
            days_in_month = 31
            header_row = 7
            total_row = header_row + days_in_month + 1
            overtime_row = total_row + 1
            signature_label_row = overtime_row + 2

            self.assertEqual(ws.cell(row=total_row, column=1).value, "Razem ilość godzin:")
            self.assertEqual(ws.cell(row=total_row, column=4).value, schedule.total_hours_for_employee(emp))
            day_total, night_total = _day_night_totals(schedule, emp)
            self.assertEqual(ws.cell(row=total_row, column=5).value, day_total)
            self.assertEqual(ws.cell(row=total_row, column=6).value, night_total)

            self.assertEqual(ws.cell(row=overtime_row, column=1).value, "Nadgodziny:")
            self.assertEqual(ws.cell(row=overtime_row, column=4).value, _overtime_str(schedule, shop, emp))

            self.assertEqual(ws.cell(row=signature_label_row, column=1).value, "Podpis pracownika:")

    def test_multiple_employees_get_one_sheet_each(self):
        schedule = MonthSchedule(2026, 8)
        emp1 = Employee(last_name="Kowalski", first_name="Adam")
        emp2 = Employee(last_name="Nowak", first_name="Ewa")
        schedule.add_employee(emp1)
        schedule.add_employee(emp2)

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/karty.xlsx"
            export_employee_cards_to_excel(schedule, 2026, 8, path, employees=[emp1, emp2])

            wb = load_workbook(path)
            self.assertEqual(len(wb.sheetnames), 2)


class ColumnWidthTests(unittest.TestCase):
    def test_day_and_night_hour_columns_are_equal_width(self):
        schedule, emp = _sample_schedule()
        exporter = _EmployeeCardImageExporter(schedule, 2026, 8, None, emp)
        col_w = exporter._column_widths()

        self.assertEqual(len(col_w), 6)
        day_w, night_w = col_w[4], col_w[5]
        # Podział nieparzystej reszty różni się co najwyżej o 1px.
        self.assertLessEqual(abs(day_w - night_w), 1)
        self.assertGreater(day_w, 150)


class ImageCardExportTests(unittest.TestCase):
    def test_export_produces_non_empty_file(self):
        schedule, emp = _sample_schedule()
        shop = ShopConfig(2026, 8)

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/karta.jpg"
            result = export_employee_card_to_image(schedule, 2026, 8, path, shop=shop, employee=emp)

            self.assertTrue(result)
            self.assertGreater(Path(path).stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
