from pathlib import Path
import math
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from PIL import Image

from export.excel_exporter import export_schedule_to_excel
from export.image_exporter import export_schedule_to_image
from model.employee import Employee
from model.month_schedule import MonthSchedule
from model.shop_config import ShopConfig

# 2026-08-03 is a Monday - a plain trade day, no Sunday/holiday edge cases.
DAY = 3


def _two_employee_schedule():
    shop = ShopConfig(2026, 8)
    shop.constraints["force_fulltime_845"] = False
    schedule = MonthSchedule(2026, 8)
    over = Employee(last_name="Ponadwymiarowy", first_name="Jan")
    under = Employee(last_name="Wramach", first_name="Ola")
    schedule.add_employee(over)
    schedule.add_employee(under)

    nominal_hours = shop.get_full_time_nominal_hours()
    days_needed = math.ceil(nominal_hours / 8) + 1
    for day in range(1, days_needed + 1):
        schedule.get_day(over, day).set_hours("08:00", "16:00")

    schedule.get_day(under, DAY).set_hours("08:00", "16:00")

    return schedule, shop, over, under


class ExcelSingleEmployeeExportTests(unittest.TestCase):
    def test_only_the_requested_employee_is_included(self):
        schedule, shop, over, under = _two_employee_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.xlsx"
            export_schedule_to_excel(schedule, 2026, 8, path, employees=[under])

            wb = load_workbook(path)
            ws = wb.active
            # Wiersz 6 to nazwisko pierwszego (i tu: jedynego) pracownika.
            self.assertEqual(ws.cell(row=6, column=1).value, under.display_name())
            # Drugi blok wierszy (9-11) nie powinien istnieć dla żadnego
            # pracownika - sprawdzamy, że tam nic nie ma.
            self.assertIsNone(ws.cell(row=9, column=1).value)

    def test_default_still_exports_every_employee(self):
        schedule, shop, over, under = _two_employee_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.xlsx"
            export_schedule_to_excel(schedule, 2026, 8, path)

            wb = load_workbook(path)
            ws = wb.active
            self.assertEqual(ws.cell(row=6, column=1).value, over.display_name())
            self.assertEqual(ws.cell(row=9, column=1).value, under.display_name())


class ExcelOvertimeHighlightTests(unittest.TestCase):
    def test_over_limit_employee_gets_red_font_when_shop_is_passed(self):
        schedule, shop, over, under = _two_employee_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.xlsx"
            export_schedule_to_excel(schedule, 2026, 8, path, shop=shop)

            wb = load_workbook(path)
            ws = wb.active
            days_in_month = 31
            sum_col_start = days_in_month + 3

            over_cell = ws.cell(row=6, column=sum_col_start)
            under_cell = ws.cell(row=9, column=sum_col_start)

        self.assertEqual(over_cell.font.color.rgb, "FFCC0000")
        self.assertNotEqual(under_cell.font.color and under_cell.font.color.rgb, "FFCC0000")

    def test_no_highlight_without_shop_argument(self):
        schedule, shop, over, under = _two_employee_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.xlsx"
            export_schedule_to_excel(schedule, 2026, 8, path)

            wb = load_workbook(path)
            ws = wb.active
            over_cell = ws.cell(row=6, column=34)

        self.assertNotEqual(over_cell.font.color and over_cell.font.color.rgb, "FFCC0000")


class ImageSingleEmployeeExportTests(unittest.TestCase):
    def test_export_with_one_employee_is_smaller_than_with_two(self):
        schedule, shop, over, under = _two_employee_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path_all = f"{tmp}/grafik_all.jpg"
            path_one = f"{tmp}/grafik_one.jpg"
            export_schedule_to_image(schedule, 2026, 8, path_all)
            export_schedule_to_image(schedule, 2026, 8, path_one, employees=[under])

            with Image.open(path_all) as img_all, Image.open(path_one) as img_one:
                self.assertGreater(img_all.height, img_one.height)

    def test_export_with_shop_does_not_crash(self):
        schedule, shop, over, under = _two_employee_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.jpg"
            result = export_schedule_to_image(schedule, 2026, 8, path, shop=shop)

            self.assertTrue(result)
            self.assertGreater(Path(path).stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
