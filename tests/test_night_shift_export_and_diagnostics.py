from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from export.excel_exporter import export_schedule_to_excel
from export.image_exporter import export_schedule_to_image
from logic.generator.diagnostics import audit_schedule
from model.constraints import ConstraintEngine, rest_11h_violation
from model.employee import Employee
from model.month_schedule import MonthSchedule
from model.shop_config import ShopConfig

# 2026-08-03/04 are a Monday/Tuesday - plain trade days, no Sunday edge cases.
DAY_1, DAY_2 = 3, 4


def _schedule_with_night_shift(gap_minutes_short=False):
    """One employee with a night shift on DAY_1 (22:00-06:00) followed by a
    DAY_2 shift. `gap_minutes_short` controls whether DAY_2's start leaves
    less than 11h of real rest after the night shift actually ends (06:00
    on DAY_2, not DAY_1)."""
    shop = ShopConfig(2026, 8)
    schedule = MonthSchedule(2026, 8)
    emp = Employee(last_name="Kowalski", first_name="Jan")
    schedule.add_employee(emp)

    schedule.get_day(emp, DAY_1).set_hours("22:00", "06:00")
    next_start = "10:00" if gap_minutes_short else "18:00"  # 4h vs 12h real rest after 06:00
    schedule.get_day(emp, DAY_2).set_hours(next_start, "20:00")

    return schedule, shop, emp


class ExcelExporterNightShiftTests(unittest.TestCase):
    def test_night_shift_cell_has_no_plus_one_marker(self):
        """Znacznik "+1" usunięty całkiem na życzenie użytkownika - tylko
        surowa godzina końca, bez oznaczenia przejścia w kolejną dobę."""
        schedule, shop, emp = _schedule_with_night_shift()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.xlsx"
            export_schedule_to_excel(schedule, 2026, 8, path)

            wb = load_workbook(path)
            ws = wb.active

            # Pierwszy pracownik zajmuje wiersze 6-8 (od/do/h), kolumna dnia
            # DAY_1 to DAY_1 + 2 (patrz excel_exporter.py: col = day + 2).
            end_cell = ws.cell(row=7, column=DAY_1 + 2).value
            start_cell = ws.cell(row=6, column=DAY_1 + 2).value
            total_cell = ws.cell(row=8, column=DAY_1 + 2).value

        self.assertEqual(start_cell, "22")
        self.assertEqual(end_cell, "6")
        self.assertNotIn("+1", end_cell)
        self.assertEqual(total_cell, "8:00")

    def test_normal_shift_cell_is_still_unaffected(self):
        schedule, shop, emp = _schedule_with_night_shift()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.xlsx"
            export_schedule_to_excel(schedule, 2026, 8, path)

            wb = load_workbook(path)
            ws = wb.active
            end_cell = ws.cell(row=7, column=DAY_2 + 2).value

        self.assertNotIn("+1", end_cell or "")


class ImageExporterNightShiftTests(unittest.TestCase):
    def test_export_with_night_shift_does_not_crash(self):
        schedule, shop, emp = _schedule_with_night_shift()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.jpg"
            result = export_schedule_to_image(schedule, 2026, 8, path)

            self.assertTrue(result)
            self.assertTrue(Path(path).exists())
            self.assertGreater(Path(path).stat().st_size, 0)


class RestElevenHourDiagnosticsNightShiftTests(unittest.TestCase):
    """model/constraints.py is the human-facing grid-coloring diagnostics -
    separate from the actual generator constraint in
    logic/generator/night_shift_constraint.py. Without accounting for
    crossing midnight, both Rest11hRule and rest_11h_violation computed the
    gap between a night shift's end and the next real day's start as if the
    night shift ended on the day it started, overestimating rest by 24h and
    never flagging a genuine violation."""

    def test_violation_detected_when_next_day_starts_too_soon_after_night_end(self):
        schedule, shop, emp = _schedule_with_night_shift(gap_minutes_short=True)

        result = rest_11h_violation(schedule, emp, DAY_1)
        self.assertTrue(result["violation"])
        self.assertAlmostEqual(result["rest_hours"], 4.0, places=1)

    def test_no_violation_when_next_day_has_enough_real_rest(self):
        schedule, shop, emp = _schedule_with_night_shift(gap_minutes_short=False)

        result = rest_11h_violation(schedule, emp, DAY_1)
        self.assertFalse(result["violation"])
        self.assertAlmostEqual(result["rest_hours"], 12.0, places=1)

    def test_constraint_engine_flags_rest_11h_violation_after_night_shift(self):
        schedule, shop, emp = _schedule_with_night_shift(gap_minutes_short=True)
        shop.constraints["enforce_11h_rest"] = True

        violations = ConstraintEngine.evaluate(schedule, shop)
        rest_violations = [v for v in violations if v.type == "rest_11h"]

        self.assertEqual(len(rest_violations), 1)
        self.assertEqual(rest_violations[0].day, DAY_2)


class GeneratorDiagnosticsNightShiftTests(unittest.TestCase):
    def test_audit_schedule_flags_short_rest_after_night_shift(self):
        schedule, shop, emp = _schedule_with_night_shift(gap_minutes_short=True)

        report = audit_schedule(schedule, shop)
        matches = [v for v in report["rest_11h_violations"] if v["days"] == f"{DAY_1}->{DAY_2}"]

        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0]["rest_minutes"], 4 * 60)

    def test_audit_schedule_does_not_flag_sufficient_rest_after_night_shift(self):
        schedule, shop, emp = _schedule_with_night_shift(gap_minutes_short=False)

        report = audit_schedule(schedule, shop)
        matches = [v for v in report["rest_11h_violations"] if v["days"] == f"{DAY_1}->{DAY_2}"]

        self.assertEqual(matches, [])


if __name__ == "__main__":
    unittest.main()
