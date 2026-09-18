from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from export.excel_exporter import export_schedule_to_excel
from export.image_exporter import export_schedule_to_image
from export.export_style import is_dino_style
from model.employee import Employee
from model.month_schedule import MonthSchedule
from model.shop_config import ShopConfig


def _one_employee_schedule():
    schedule = MonthSchedule(2026, 8)
    emp = Employee(last_name="Testowy", first_name="Jan")
    schedule.add_employee(emp)
    schedule.get_day(emp, 3).set_hours("08:00", "16:00")
    return schedule, emp


class IsDinoStyleTests(unittest.TestCase):
    def test_none_shop_counts_as_dino(self):
        self.assertTrue(is_dino_style(None))

    def test_dino_retail_business_type_is_dino(self):
        shop = ShopConfig(2026, 8)
        self.assertTrue(is_dino_style(shop))

    def test_other_business_type_is_not_dino(self):
        shop = ShopConfig(2026, 8)
        shop.business_type = "ochrona_enyo"
        self.assertFalse(is_dino_style(shop))


class ExcelDispatchTests(unittest.TestCase):
    def test_dino_profile_keeps_todays_layout(self):
        schedule, emp = _one_employee_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.xlsx"
            export_schedule_to_excel(schedule, 2026, 8, path)

            wb = load_workbook(path)
            ws = wb.active
            header_values = [ws.cell(row=r, column=c).value for r in range(1, 4) for c in range(1, ws.max_column + 1)]
            self.assertIn("Wydruk wewnętrzny", header_values)

    def test_non_dino_profile_uses_security_exporter(self):
        schedule, emp = _one_employee_schedule()
        shop = ShopConfig(2026, 8)
        shop.business_type = "ochrona_enyo"

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.xlsx"

            with patch(
                "export.security_excel_exporter.export_security_schedule_to_excel",
                return_value=None,
            ) as mocked:
                export_schedule_to_excel(schedule, 2026, 8, path, shop=shop)
                self.assertTrue(mocked.called)

    def test_non_dino_profile_output_has_no_dino_branding(self):
        schedule, emp = _one_employee_schedule()
        shop = ShopConfig(2026, 8)
        shop.business_type = "ochrona_enyo"

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.xlsx"
            export_schedule_to_excel(schedule, 2026, 8, path, shop=shop)

            wb = load_workbook(path)
            ws = wb.active
            all_values = [
                ws.cell(row=r, column=c).value
                for r in range(1, 6)
                for c in range(1, 6)
            ]
            self.assertNotIn("Wydruk wewnętrzny", all_values)


class ImageDispatchTests(unittest.TestCase):
    def test_non_dino_profile_uses_security_exporter(self):
        schedule, emp = _one_employee_schedule()
        shop = ShopConfig(2026, 8)
        shop.business_type = "ochrona_enyo"

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.jpg"

            with patch(
                "export.security_image_exporter.export_security_schedule_to_image",
                return_value=True,
            ) as mocked:
                export_schedule_to_image(schedule, 2026, 8, path, shop=shop)
                self.assertTrue(mocked.called)

    def test_dino_profile_still_produces_output(self):
        schedule, emp = _one_employee_schedule()

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.jpg"
            result = export_schedule_to_image(schedule, 2026, 8, path)
            self.assertTrue(result)
            self.assertGreater(Path(path).stat().st_size, 0)

    def test_non_dino_profile_produces_real_output_without_mocking(self):
        schedule, emp = _one_employee_schedule()
        shop = ShopConfig(2026, 8)
        shop.business_type = "ochrona_enyo"

        with tempfile.TemporaryDirectory() as tmp:
            path = f"{tmp}/grafik.jpg"
            result = export_schedule_to_image(schedule, 2026, 8, path, shop=shop)
            self.assertTrue(result)
            self.assertGreater(Path(path).stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
